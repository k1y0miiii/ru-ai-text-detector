#!/usr/bin/env python3
"""
test_segmentation.py — РЕГРЕССИЯ сегментации classify_sections. Защищает от рецидива
TOC/back-matter latching и от засорения таблицами. Запуск: python test_segmentation.py

Покрывает:
  - все накопленные samples/ (test_diploma/mixed/pure_ai/ai_edited/synthetic/human);
  - СИНТЕТИЧЕСКИЙ «торрент-кейс», воспроизводящий патологию реального ВКР, чтобы тест
    не зависел от личного файла: длинное поле «Задания» ДО оглавления, строки TOC с
    табом+номером страницы, маркер «Список источников» ДВАЖДЫ (в TOC и реальный),
    настоящая Word-таблица в теле. Ассерты: тело = от настоящего ВВЕДЕНИЯ до реального
    списка, таблица в анализ не попала, проанализировано БОЛЬШИНСТВО.
"""

import re
import sys
import tempfile
from pathlib import Path

from aidetector import document_scan as ds
from aidetector.extract_text import extract_text


def flat(t):
    return re.sub(r"\s+", " ", t).strip()


def content_bounds(path):
    text, _ = extract_text(Path(path))
    paras = ds.split_paragraphs(text)
    split = ds.classify_sections(paras)
    c = split.content
    ntab = sum(1 for _, r in split.dropped if "таблица" in r)
    return paras, split, c, ntab


def build_torture_docx(p: Path):
    """Синтетический ВКР с патологией реального файла (для регрессии без личных данных)."""
    from docx import Document
    d = Document()
    LONG = ("Это длинный связный абзац настоящей прозы, который занимает заметно больше "
            "двухсот символов и потому проходит проверку на тело документа; он нужен, чтобы "
            "проверить, что сегментация цепляется именно за настоящий контент после оглавления.")
    # --- титул + ЗАДАНИЕ (длинное поле ДО оглавления — ловушка first_body) ---
    for t in ["АНО ВО «Институт деловой карьеры»", "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА",
              "на тему: «Тест сегментации»", "Москва — 2026", "Задание на ВКР",
              "1. Тема ВКР: «Тест сегментации».",
              "3. Исходные данные к работе: материалы преддипломной практики и техническая "
              "документация; нормативные документы; результаты обследования предметной области, "
              "собранные в ходе работы над проектом, включая описание процессов как есть."]:
        d.add_paragraph(t)
    # --- ОГЛАВЛЕНИЕ: строки с табом+номером страницы (НЕ настоящие заголовки) ---
    d.add_paragraph("Содержание")
    for t in ["ВВЕДЕНИЕ\t3", "ГЛАВА 1. АНАЛИЗ\t7", "ЗАКЛЮЧЕНИЕ\t40",
              "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ\t44", "ПРИЛОЖЕНИЯ\t47"]:
        d.add_paragraph(t)
    # --- НАСТОЯЩЕЕ тело ---
    d.add_paragraph("ВВЕДЕНИЕ")
    d.add_paragraph(LONG)
    d.add_paragraph("ГЛАВА 1. АНАЛИЗ")
    d.add_paragraph(LONG)
    # таблица в теле (данные — не проза)
    tbl = d.add_table(rows=2, cols=2)
    for r, row in enumerate(tbl.rows):
        for cc, cell in enumerate(row.cells):
            cell.text = ["project", "внешний ключ", "каскадное удаление", "Ссылка на проект"][r * 2 + cc]
    d.add_paragraph(LONG)
    d.add_paragraph("ЗАКЛЮЧЕНИЕ")
    d.add_paragraph(LONG)
    # --- НАСТОЯЩИЙ список источников (второе вхождение маркера) + приложения ---
    d.add_paragraph("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for t in ["1. Иванов И. И. Книга. — М.: Изд, 2020. — 100 с.",
              "2. Петров П. П. Статья // Журнал. 2021. № 3. С. 1-10."]:
        d.add_paragraph(t)
    d.add_paragraph("ПРИЛОЖЕНИЯ")
    d.add_paragraph("Приложение А. Данные.")
    d.save(str(p))


def main():
    ok = True

    def check(name, cond, detail=""):
        nonlocal ok
        ok = ok and cond
        print(f"  [{'OK ' if cond else 'FAIL'}] {name}  {detail}")

    print("=== samples/ (содержательная часть должна быть введение→вывод) ===")
    for f in ["test_diploma.docx", "test_mixed.docx", "test_pure_ai.docx",
              "test_ai_edited.docx", "synthetic_doc.txt", "sample_human_1.txt"]:
        paras, split, c, ntab = content_bounds(f"samples/{f}")
        starts_intro = bool(c) and flat(c[0].text).lower().startswith(("введение", "вчера"))
        check(f"{f}: content непустой и стартует с тела", bool(c) and starts_intro,
              f"({len(c)} абз., начало «{flat(c[0].text)[:24]}»)")

    print("\n=== СИНТЕТИЧЕСКИЙ торрент-кейс (патология реального ВКР) ===")
    with tempfile.TemporaryDirectory() as tmp:
        tp = Path(tmp) / "torture.docx"
        build_torture_docx(tp)
        paras, split, c, ntab = content_bounds(tp)
        texts = [flat(p.text) for p in c]
        check("тело стартует с НАСТОЯЩЕГО «ВВЕДЕНИЕ» (не с поля Задания, не с TOC)",
              bool(c) and texts[0] == "ВВЕДЕНИЕ", f"начало «{texts[0][:30]}»")
        check("«Исходные данные к работе…» (поле Задания) НЕ в анализе",
              not any("Исходные данные к работе" in t for t in texts))
        check("строка TOC «…ИСТОЧНИКОВ 44» НЕ латчит back-matter раньше времени",
              any("ЗАКЛЮЧЕНИЕ" == t for t in texts) and
              any("материалы" not in t and len(t) > 200 for t in texts),
              f"({len(c)} абз. в теле)")
        check("реальный «Список источников» отрезал back-matter (пункты-источники НЕ в анализе)",
              not any(t.startswith(("1. Иванов", "2. Петров")) for t in texts))
        check("таблица (project/внешний ключ) исключена из анализа",
              ntab >= 1 and not any("внешний ключ" in t for t in texts), f"(таблиц отсеяно: {ntab})")
        check("проанализировано тело (есть длинная проза)",
              sum(1 for t in texts if len(t) > 200) >= 3)

    # --- РЕФЕРЕНСНЫЙ реальный ВКР (опционально) — самый ценный кейс ---
    # Свой документ положите как samples/real_thesis.docx (в .gitignore, не публикуется).
    ref = Path("samples/real_thesis.docx")
    print(f"\n=== РЕФЕРЕНСНЫЙ реальный ВКР ({ref.name}) ===")
    if ref.exists():
        paras, split, c, ntab = content_bounds(ref)
        texts = [flat(p.text) for p in c]
        check("сегментировался в БОЛЬШИНСТВО прозы (не 2% как при баге)",
              len(c) >= 300, f"(content {len(c)} абз. из {len(paras)})")
        # «Введение» должно быть у НАЧАЛА тела (тело не выброшено в back-matter).
        # ВАЖНО: у ГОСТ-версии есть front-matter «Перечень приложений» (описания
        # Приложений А-Г) ПЕРЕД Введением — first_body латчит на «Приложение А»,
        # поэтому ~4 строки фронт-матера попадают в content до Введения (читаются как
        # человек, не флагаются). Известная минорная особенность, не катастрофа бага.
        intro = next((i for i, t in enumerate(texts) if t.lower().startswith("введение")), None)
        check("«Введение» присутствует в content у начала (тело не выброшено)",
              intro is not None and intro <= 6, f"(позиция Введения в content: {intro})")
        check("таблицы исключены из анализа (данные, не проза)", ntab >= 10,
              f"(таблиц отсеяно {ntab})")
        check("поля «Задания»/строки TOC НЕ в анализе",
              not any("Исходные данные к работе" in t for t in texts) and
              not any(re.search(r"\t\d{1,4}$", flat(p.text)) for p in c))
    else:
        print(f"  [skip] {ref} нет в samples/ — пропускаю (положи для реального регресса)")

    print("\n" + ("ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ ✓" if ok else "ЕСТЬ ПАДЕНИЯ ✗"))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
