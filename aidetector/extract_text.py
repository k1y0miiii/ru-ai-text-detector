#!/usr/bin/env python3
"""
extract_text.py — извлечение ЧИСТОГО текста из .docx / .pdf / .txt для подачи в
document_scan.py. Это препроцессинг ПЕРЕД детектором: модель/порог/features не
трогаем, тут только «достать авторский текст из файла».

Главный контракт с document_scan.split_paragraphs (он режет по пустым строкам):
  * абзацы на выходе разделены ПУСТОЙ строкой ('\\n\\n');
  * заголовки («Введение», «Глава 1», «Список литературы») остаются ОТДЕЛЬНЫМИ
    абзацами — иначе сломается отсев служебных секций по белому списку.

Что НЕ извлекаем:
  * картинки/иллюстрации/графики/скрины — это не авторский текст, в анализ не идут;
  * OCR сканов НЕ делаем (отдельный этап). Если текста-как-текста почти нет, а
    страницы/картинки есть — честно предупреждаем «похоже на скан», а не молчим.
"""

import re
import sys
import zlib
from pathlib import Path

# Ниже этого числа символов считаем, что текста-как-текста в файле нет
# (вероятно скан или картинка вместо текста).
SCAN_TEXT_THRESHOLD = 100

# Маркер таблицы. Таблицы — это ДАННЫЕ (схемы БД, расчёты, перечни ячеек вроде
# «project», «внешний ключ»), а НЕ авторская проза. Разворачивать их в мелкие абзацы-
# ячейки нельзя — они засоряют окна анализа. Заменяем каждую таблицу на компактный
# служебный маркер; document_scan его распознаёт и в анализ НЕ берёт (но показывает,
# что таблица была — не маскируем). Символы ⟦⟧ почти не встречаются в тексте.
TABLE_MARK = "⟦таблица"


def _clean_para(s: str) -> str:
    """Внутренние переносы строк -> пробел: абзац остаётся ОДНИМ абзацем."""
    return re.sub(r"[ \t]*\n[ \t]*", " ", s).strip()


# =============================================================================
# DOCX
# =============================================================================
def _iter_block_items(parent):
    """
    Идём по содержимому docx В ПОРЯДКЕ следования, отдавая параграфы и таблицы.
    python-docx по отдельности знает .paragraphs и .tables, но теряет ИХ ПОРЯДОК
    относительно друг друга — поэтому обходим XML-тело напрямую.
    """
    from docx.document import Document as _Doc
    from docx.oxml.ns import qn
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        parent_elm = parent

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _docx_table_paragraphs(table) -> list[str]:
    """Текст таблицы по ячейкам (в них часто содержательный текст). Пусто — пропуск."""
    out = []
    for row in table.rows:
        for cell in row.cells:
            parts = []
            for item in _iter_block_items(cell):
                # ячейка может содержать вложенные таблицы
                if item.__class__.__name__ == "Table":
                    parts.extend(_docx_table_paragraphs(item))
                else:
                    t = _clean_para(item.text)
                    if t:
                        parts.append(t)
            cell_text = " ".join(parts).strip()
            if cell_text:
                out.append(cell_text)
    return out


# Подпись к рисунку/таблице по ТЕКСТУ («Рисунок 1.1 — …», «Таблица 2 …»).
_CAPTION_TEXT_RE = re.compile(
    r"^\s*(рис(?:унок)?\.?|табл(?:ица)?\.?|схема|диаграмма|листинг|график)\s*№?\s*\d",
    re.IGNORECASE)


def _para_xml_meta(block) -> dict:
    """XML-признаки параграфа docx для УМНОЙ категоризации коротких абзацев:
    стиль (Heading/Caption), наличие Office Math (m:oMath), подпись по тексту.
    Только ЧИТАЕМ разметку — модель/сегментацию не трогаем."""
    style = ""
    try:
        style = (block.style.name or "").strip().lower()
    except Exception:
        style = ""
    is_heading = "heading" in style or "заголовок" in style or style == "title" or "название" in style
    is_caption_style = "caption" in style or "подпись" in style or "название объекта" in style
    txt = (block.text or "").strip()
    is_caption = is_caption_style or bool(_CAPTION_TEXT_RE.match(txt))
    # Office Math: формула как объект OMML (m:oMath) внутри параграфа
    has_math = False
    try:
        has_math = "oMath" in block._p.xml
    except Exception:
        has_math = False
    return {"style": style, "is_heading": is_heading,
            "is_caption": is_caption, "has_math": has_math, "is_table": False}


def _extract_docx(path: Path):
    """-> (paragraphs: list[str], warnings: list[str], blocks: list[dict]).
    blocks выровнены 1:1 с paragraphs (каждый emitted-абзац -> один dict XML-мета);
    таблицы-маркеры тоже представлены (is_table=True)."""
    from docx import Document

    doc = Document(str(path))
    paras: list[str] = []
    blocks: list[dict] = []
    n_tables = 0
    for block in _iter_block_items(doc):
        if block.__class__.__name__ == "Table":
            cells = _docx_table_paragraphs(block)
            if cells:
                # таблицу в анализ НЕ берём (данные, не проза) — компактный маркер
                n_tables += 1
                paras.append(f"{TABLE_MARK}: {len(cells)} ячеек⟧")
                blocks.append({"style": "table", "is_heading": False,
                               "is_caption": False, "has_math": False, "is_table": True})
        else:
            # paragraph.text возвращает ТОЛЬКО текстовые run'ы — рисунки/картинки
            # (w:drawing) в .text не попадают, поэтому изображения игнорируются сами.
            t = _clean_para(block.text)
            if t:
                paras.append(t)
                blocks.append(_para_xml_meta(block))

    warnings: list[str] = []
    total = sum(len(p) for p in paras)
    n_images = len(getattr(doc, "inline_shapes", []))
    if total < SCAN_TEXT_THRESHOLD and n_images > 0:
        warnings.append(
            f"похоже на скан: текста-как-текста почти нет ({total} симв), "
            f"а картинок {n_images} — нужен OCR (отдельный этап, пока не реализован)"
        )
    return paras, warnings, blocks


# =============================================================================
# PDF (PyMuPDF / fitz)
# =============================================================================
def _extract_pdf(path: Path):
    """-> (paragraphs: list[str], warnings: list[str]). Картинки-блоки пропускаем."""
    import fitz

    paras: list[str] = []
    n_images = 0
    with fitz.open(str(path)) as doc:
        n_pages = doc.page_count
        for page in doc:
            n_images += len(page.get_images())
            # "blocks": (x0, y0, x1, y1, text, block_no, block_type), type 1 = картинка
            blocks = page.get_text("blocks")
            blocks.sort(key=lambda b: (round(b[1], 1), b[0]))  # порядок чтения
            for b in blocks:
                if len(b) >= 7 and b[6] != 0:
                    continue  # блок-картинка — игнорируем
                t = _clean_para(b[4])
                if t:
                    paras.append(t)

    warnings: list[str] = []
    total = sum(len(p) for p in paras)
    if total < SCAN_TEXT_THRESHOLD and (n_pages > 0 or n_images > 0):
        warnings.append(
            f"похоже на скан: текста-как-текста почти нет ({total} симв) при "
            f"{n_pages} стр./{n_images} картинках — нужен OCR "
            "(отдельный этап, пока не реализован)"
        )
    return paras, warnings


# =============================================================================
# TXT
# =============================================================================
def _extract_txt(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    # .txt уже в нужном формате (абзацы пустой строкой) — отдаём как есть
    return text, []


# =============================================================================
# Публичный интерфейс
# =============================================================================
def extract_blocks(path) -> tuple[str, list[str], "list[dict] | None"]:
    """
    Как extract_text, но ДОПОЛНИТЕЛЬНО возвращает per-абзац XML-метаданные (blocks).

    -> (text, warnings, blocks):
      blocks — для .docx: список dict {style,is_heading,is_caption,has_math,is_table},
               ВЫРОВНЕННЫЙ 1:1 с абзацами после split_paragraphs (тот же порядок/индекс);
               для .txt/.pdf — None (XML-разметки нет, категоризация падает на эвристику).
    """
    path = Path(path)
    if not path.exists():
        return "", [f"файл не найден: {path}"], None

    ext = path.suffix.lower()
    if ext == ".txt":
        text, warnings = _extract_txt(path)
        return text, warnings, None
    elif ext == ".docx":
        paras, warnings, blocks = _extract_docx(path)
    elif ext == ".pdf":
        paras, warnings = _extract_pdf(path)
        blocks = None
    elif ext == ".doc":
        return "", [".doc (старый формат) не поддерживается — пересохрани в .docx"], None
    else:
        return "", [f"неизвестный формат «{ext}» — поддерживаются .docx/.pdf/.txt"], None

    text = "\n\n".join(paras)
    if not text.strip() and not warnings:
        warnings.append("из файла не извлёкся текст (пустой документ?)")
    return text, warnings, blocks


def extract_text(path) -> tuple[str, list[str]]:
    """
    Извлечь чистый текст из файла. Формат — по расширению (.docx/.pdf/.txt).

    Возвращает (text, warnings):
      text     — абзацы, разделённые пустой строкой (готово для document_scan);
      warnings — список предупреждений (например, «похоже на скан»). Пустой список,
                 если всё чисто. Пустоту НЕ маскируем — при подозрении на скан
                 предупреждение обязательно присутствует.
    (Обёртка над extract_blocks — обратная совместимость со всеми вызовами.)
    """
    text, warnings, _ = extract_blocks(path)
    return text, warnings


def main():
    if len(sys.argv) != 2:
        sys.exit("использование: python extract_text.py <файл.docx|.pdf|.txt>")
    text, warnings = extract_text(sys.argv[1])
    for w in warnings:
        print(f"[!] {w}", file=sys.stderr)
    if text:
        print(text)


if __name__ == "__main__":
    main()
