#!/usr/bin/env python3
"""
build_test_docx.py — собирает РЕАЛЬНЫЙ тестовый .docx для проверки полного пути
extract_text -> document_scan. Структура как у настоящего диплома: титульник,
содержание, Введение, две главы (с длинными ОДНОЗНАЧНЫМИ AI- и human-абзацами из
data/test.jsonl — как в синтетике), Заключение, Список литературы, Приложение с
таблицей, плюс картинка-заглушка (должна игнорироваться при извлечении).

Картинку-PNG генерим стандартным zlib (без Pillow — лишняя зависимость не нужна).
"""

import json
import re
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.shared import Inches

S = Path("samples")


def make_png(path: Path, w=160, h=90, color=(200, 200, 205)) -> None:
    """Минимальный валидный RGB-PNG сплошного цвета, чисто на stdlib."""
    def chunk(typ: bytes, data: bytes) -> bytes:
        body = typ + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    raw = bytearray()
    row = bytes(color) * w
    for _ in range(h):
        raw.append(0)          # фильтр строки = 0
        raw.extend(row)
    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))  # RGB, 8 бит
           + chunk(b"IDAT", zlib.compress(bytes(raw)))
           + chunk(b"IEND", b""))
    path.write_bytes(png)


# --- те же длинные РАЗНЫЕ размеченные абзацы, что в синтетике ---------------
rows = [json.loads(l) for l in open("data/test.jsonl", encoding="utf-8")]
wc = lambda t: len(t.split())
flat = lambda t: re.sub(r"\s+", " ", t).strip()
ai = sorted((x for x in rows if x["label"] == 1), key=lambda x: -wc(x["text"]))
hum = sorted((x for x in rows if x["label"] == 0), key=lambda x: -wc(x["text"]))
AI1, AI2 = flat(ai[1]["text"]), flat(ai[3]["text"])      # p~1.0
HUM1, HUM2 = flat(hum[1]["text"]), flat(hum[4]["text"])  # p~0.05

conclusion = (
    "В общем, пока возился со всем этим, понял простую вещь: проще сразу делать "
    "по-нормальному, чем потом переделывать. Половину времени убил на ерунду, "
    "которую можно было обойти, если бы не торопился. Зато теперь хоть знаю, как "
    "надо, и в следующий раз точно не полезу наобум — обидно за потраченные выходные."
)

doc = Document()

# --- титульник ---
for line in [
    "Министерство образования и науки Российской Федерации",
    "Федеральный университет, кафедра прикладной информатики",
    "КУРСОВАЯ РАБОТА",
    "на тему: «Применение методов анализа текста»",
    "Выполнил: студент группы ПИ-301",
    "Проверил: доц. Иванов И.И.",
    "Москва — 2024",
]:
    doc.add_paragraph(line)

# --- оглавление (служебное) ---
doc.add_heading("Содержание", level=1)
doc.add_paragraph(
    "Введение .......... 3\nГлава 1 .......... 5\nГлава 2 .......... 12\n"
    "Заключение .......... 20\nСписок литературы .......... 22")

# --- содержательная часть: AI-регион ---
doc.add_heading("Введение", level=1)
doc.add_paragraph(AI1)
doc.add_heading("Глава 1. Теоретическая часть", level=1)
doc.add_paragraph(AI2)

# картинка-заглушка + её подпись (картинка ИГНОРИРУЕТСЯ, подпись — текст автора)
fig = S / "_fig.png"
make_png(fig)
doc.add_picture(str(fig), width=Inches(2.5))
doc.add_paragraph("Рисунок 1 — иллюстрация к главе")

# --- содержательная часть: human-регион ---
doc.add_heading("Глава 2. Практическая часть", level=1)
doc.add_paragraph(HUM1)
doc.add_paragraph(HUM2)

doc.add_heading("Заключение", level=1)
doc.add_paragraph(conclusion)

# --- back-matter (служебное) ---
doc.add_heading("Список литературы", level=1)
doc.add_paragraph("1. Петров П.П. Анализ текстов. — М.: Наука, 2020. — 240 с.")
doc.add_paragraph("2. Сидоров С.С. Машинное обучение. — СПб.: Питер, 2021. — 320 с.")

doc.add_heading("Приложение А", level=1)
# таблица с содержательным текстом — проверка извлечения таблиц
tbl = doc.add_table(rows=2, cols=2)
tbl.rows[0].cells[0].text = "Параметр"
tbl.rows[0].cells[1].text = "Значение"
tbl.rows[1].cells[0].text = "Объём тестовой выборки"
tbl.rows[1].cells[1].text = "260 размеченных текстов из тестового набора"

out = S / "test_diploma.docx"
doc.save(str(out))
print(f"saved {out}  (картинка-заглушка: {fig}, 1 шт.)")
